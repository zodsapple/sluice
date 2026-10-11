#!/usr/bin/env python3
"""Author the holdout cases: template.json, validation.json and expected.csv for each downloaded public file.

Independent of sluice: only csv / openpyxl / datetime are used, and every expected value is computed from the cells by the
rules written in that case's template ("rules"). Run:  py -I tools/build_holdout.py
"""
from __future__ import annotations

import calendar
import csv
import json
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
CASES = ROOT / "cases"
TOL = "0.000001"


def write_case(cid, fields, example, rules, rows, key, scope, validation, raw_file, domain, fmt):
    """fields: [(name, type, nullable)]; rows: list of dicts with every field (missing -> '')."""
    d = CASES / cid
    names = [f[0] for f in fields]
    rows = [{k: (v.strip(" \t\r\n") if isinstance(v, str) else v) for k, v in r.items()} for r in rows]      # convention for every case
    rules = [*rules, "Leading and trailing spaces of text are removed."]
    example = {k: (v.strip(" \t\r\n") if isinstance(v, str) else v) for k, v in example.items()}
    keys = {tuple(r[k] for k in key) for r in rows}
    assert len(keys) == len(rows), f"{cid}: evaluation key is not unique ({len(keys)} of {len(rows)})"
    with (d / "expected.csv").open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=names)
        w.writeheader()
        for r in rows:
            w.writerow({n: ("" if r.get(n) is None else r[n]) for n in names})
    (d / "template.json").write_text(json.dumps({
        "schema_name": f"{cid}_v1",
        "fields": [{"name": n, "type": t, "nullable": nl} for n, t, nl in fields],
        "example_rows": [example], "null_csv_representation": "empty field", "rules": rules}, indent=1, ensure_ascii=False), encoding="utf-8")
    (d / "validation.json").write_text(json.dumps(validation, indent=1, ensure_ascii=False), encoding="utf-8")
    numeric = {n: TOL for n, t, _ in fields if t in ("integer", "decimal")}
    return {"case": {"case_id": cid, "domain": domain, "format": fmt, "raw_file": f"cases/{cid}/{raw_file}", "expected_rows": len(rows)},
            "scoring": {"expected": f"cases/{cid}/expected.csv", "expected_rows": len(rows), "key": key,
                        "compare_fields": [n for n in names if n not in key and n != "source_id"], "numeric_fields": numeric,
                        "sum_controls": [], "scope": scope}}


def num(x):
    return "" if x in (None, "") else (str(int(x)) if isinstance(x, float) and x.is_integer() else str(x))


def month_end(y, m):
    return date(y, m, calendar.monthrange(y, m)[1])


# ---- 1. World Bank: wide years -> one row per country and year ------------------------------------------------------------
def wb():
    cid = "wb_population"
    recs = list(csv.reader(open(CASES / cid / "API_SP.POP.TOTL_DS2_en_csv_v2_5254.csv", encoding="utf-8-sig", newline="")))
    hi = next(i for i, r in enumerate(recs) if r and r[0] == "Country Name")
    years = recs[hi][4:]
    rows = []
    for r in recs[hi + 1:]:
        if not r:
            continue
        for y, v in zip(years, r[4:]):
            if y.strip().isdigit() and v.strip() != "":
                rows.append({"source_id": cid, "country_code": r[1], "country_name": r[0], "year": y, "population": v.strip()})
    f = [("source_id", "string", False), ("country_code", "string", False), ("country_name", "string", False),
         ("year", "integer", False), ("population", "integer", False)]
    ex = dict(rows[0])
    return write_case(cid, f, ex, ["One row per country (or aggregate) and year.", "Years without a published population are left out."],
                      rows, ["country_code", "year"], "All countries and aggregates, every year with a value",
                      {"file": "API_SP.POP.TOTL_DS2_en_csv_v2_5254.csv", "header_record": hi + 1}, "API_SP.POP.TOTL_DS2_en_csv_v2_5254.csv", "demography", "csv")


# ---- 2. ONS CPI: annual, quarterly and monthly rows in one column --------------------------------------------------------
def ons():
    cid = "ons_cpi"
    rows = []
    MONTHS = {m.upper(): i for i, m in enumerate(calendar.month_abbr) if m}
    for r in csv.reader(open(CASES / cid / "ons_cpi_d7bt.csv", encoding="utf-8-sig", newline="")):
        if len(r) < 2 or not re.match(r"^\d{4}", r[0]):
            continue
        p = r[0].strip()
        if re.fullmatch(r"\d{4}", p):
            t, s, e = "year", date(int(p), 1, 1), date(int(p), 12, 31)
        elif re.fullmatch(r"\d{4} Q[1-4]", p):
            y, q = int(p[:4]), int(p[-1])
            t, s, e = "quarter", date(y, 3 * q - 2, 1), month_end(y, 3 * q)
        else:
            y, m = int(p[:4]), MONTHS[p[5:8]]
            t, s, e = "month", date(y, m, 1), month_end(y, m)
        rows.append({"source_id": cid, "period": p, "period_type": t, "period_start": s.isoformat(), "period_end": e.isoformat(),
                     "index_value": r[1].strip()})
    f = [("source_id", "string", False), ("period", "string", False), ("period_type", "string", False), ("period_start", "date", False),
         ("period_end", "date", False), ("index_value", "decimal", False)]
    ex = next(r for r in rows)
    return write_case(cid, f, ex, ["Annual, quarterly and monthly values are all kept; period_type is year, quarter or month.",
                                   "period is the label as printed; period_start / period_end are the first and last day of the period."],
                      rows, ["period_type", "period"], "All annual, quarterly and monthly rows", {"file": "ons_cpi_d7bt.csv"}, "ons_cpi_d7bt.csv", "prices", "csv")


# ---- 3 / 4. Census population tables ------------------------------------------------------------------------------------
def census_rows(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    out = []
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < 5 or r[0] is None or not isinstance(r[1], (int, float)):
            continue
        out.append((i, str(r[0]), r[1:]))
    return out


def census(cid, fname, level_of, name_of):
    rows = []
    for i, label, vals in census_rows(CASES / cid / fname):
        lv = level_of(label)
        if lv is None:
            continue
        base, *est = vals
        for kind, ref, v in [("base", "2020-04-01", base), *[("estimate", f"{y}-07-01", v) for y, v in zip(range(2020, 2024), est)]]:
            rows.append({"source_id": cid, **name_of(label), **({"geography_level": lv} if cid == "census_state_pop" else {}),
                         "reference_date": ref, "estimate_type": kind, "population": num(v)})
    return rows


def census_state():
    cid = "census_state_pop"
    regions = {"Northeast", "Midwest", "South", "West"}

    def level(label):
        if label == "United States":
            return "nation"
        if label in regions:
            return "region"
        return "territory" if label == ".Puerto Rico" else "state"

    rows = census(cid, "NST-EST2023-POP.xlsx", level, lambda l: {"geography": l.lstrip(".")})
    f = [("source_id", "string", False), ("geography", "string", False), ("geography_level", "string", False),
         ("reference_date", "date", False), ("estimate_type", "string", False), ("population", "integer", False)]
    return write_case(cid, f, rows[0], ["The leading dot of a geography is layout, not part of its name.",
                                        "geography_level is nation, region, state or territory; the District of Columbia counts as a state, Puerto Rico as a territory.",
                                        "The April 1, 2020 column is the estimates base (estimate_type base, reference_date 2020-04-01); the other columns are July 1 estimates.",
                                        "Notes and citations under the table are not data."],
                      rows, ["geography", "reference_date", "estimate_type"], "Every geography row, base and four annual estimates",
                      {"sheet": "NST-EST2023-POP"}, "NST-EST2023-POP.xlsx", "demography", "xlsx")


def census_county():
    cid = "census_county_pop"

    def level(label):
        return "county" if label.startswith(".") else None

    def names(label):
        county, state = label.lstrip(".").split(", ", 1)
        return {"county": county, "state": state}

    rows = census(cid, "co-est2023-pop-01.xlsx", level, names)
    f = [("source_id", "string", False), ("county", "string", False), ("state", "string", False), ("reference_date", "date", False),
         ("estimate_type", "string", False), ("population", "integer", False)]
    return write_case(cid, f, rows[0], ["Only county rows; the state total row is not a county.",
                                        "The row label reads '.Autauga County, Alabama': county is 'Autauga County', state is 'Alabama'.",
                                        "The April 1, 2020 column is the estimates base (estimate_type base, reference_date 2020-04-01); the other columns are July 1 estimates."],
                      rows, ["county", "state", "reference_date", "estimate_type"], "Every county row, base and four annual estimates",
                      {"sheet": "CO-EST2023-POP-01"}, "co-est2023-pop-01.xlsx", "demography", "xlsx")


# ---- 5. Eurostat TSV: composite key columns, flags next to values ---------------------------------------------------------
def eurostat():
    cid = "eurostat_unemployment"
    lines = [l.split("\t") for l in open(CASES / cid / "estat_une_rt_a.tsv", encoding="utf-8").read().splitlines()]
    years = [y.strip() for y in lines[0][1:]]
    rows = []
    for r in lines[1:]:
        _, age, unit, sex, geo = r[0].split(",")
        for y, cell in zip(years, r[1:]):
            c = cell.strip()
            if c.startswith(":") or c == "":
                continue
            m = re.fullmatch(r"(-?[\d.]+)\s*([a-z]*)", c)
            rows.append({"source_id": cid, "geo": geo, "age_group": age, "sex": sex, "unit": unit, "year": y, "rate": m.group(1),
                         "flag": m.group(2)})
    f = [("source_id", "string", False), ("geo", "string", False), ("age_group", "string", False), ("sex", "string", False),
         ("unit", "string", False), ("year", "integer", False), ("rate", "decimal", False), ("flag", "string", True)]
    return write_case(cid, f, rows[0], ["Missing values (written ':') are left out.", "A letter after a value is a Eurostat flag: the number goes to rate, the letter to flag."],
                      rows, ["geo", "age_group", "sex", "unit", "year"], "Every series and year with a value", {"file": "estat_une_rt_a.tsv"},
                      "estat_une_rt_a.tsv", "labour", "tsv")


# ---- 6. CDC: long file, a human states the scope ---------------------------------------------------------------------------
def cdc():
    cid = "cdc_covid_deaths"
    rows = []
    for r in csv.DictReader(open(CASES / cid / "Provisional_COVID-19_Deaths_by_Sex_and_Age.csv", encoding="utf-8-sig", newline="")):
        if r["Group"] != "By Year":
            continue
        d = datetime.strptime(r["Data As Of"], "%m/%d/%Y").date().isoformat()
        rows.append({"source_id": cid, "data_as_of": d, "year": r["Year"], "state": r["State"], "sex": r["Sex"], "age_group": r["Age Group"],
                     "covid_deaths": r["COVID-19 Deaths"], "total_deaths": r["Total Deaths"]})
    f = [("source_id", "string", False), ("data_as_of", "date", False), ("year", "integer", False), ("state", "string", False),
         ("sex", "string", False), ("age_group", "string", False), ("covid_deaths", "integer", True), ("total_deaths", "integer", True)]
    return write_case(cid, f, rows[0], ["Only the rows whose Group is 'By Year'.", "Dates are printed month/day/year.",
                                        "A blank count (suppressed) stays empty, never zero."],
                      rows, ["year", "state", "sex", "age_group"], "Rows with Group = By Year", {"file": "Provisional_COVID-19_Deaths_by_Sex_and_Age.csv",
                                                                                                "filter": {"Group": "By Year"}},
                      "Provisional_COVID-19_Deaths_by_Sex_and_Age.csv", "health", "csv")


# ---- 7. UCI online retail: 540k transaction lines ---------------------------------------------------------------------------
def retail():
    cid = "uci_online_retail"
    ws = openpyxl.load_workbook(CASES / cid / "Online Retail.xlsx", read_only=True, data_only=True).active
    ws.reset_dimensions()
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == 1:
            continue
        inv, stock, desc, qty, when, price, cust, country = r[:8]
        if inv is None:
            continue
        q, p = int(qty), Decimal(str(price))
        rows.append({"source_id": cid, "invoice_no": str(inv), "stock_code": str(stock), "description": "" if desc is None else str(desc),
                     "quantity": str(q), "invoice_date": when.date().isoformat(), "unit_price": str(p),
                     "customer_id": "" if cust is None else str(int(cust)), "country": country,
                     "is_cancellation": "true" if str(inv).startswith("C") else "false", "line_amount": str((p * q).quantize(Decimal("0.0001"))),
                     "source_row": str(i)})
    f = [("source_id", "string", False), ("invoice_no", "string", False), ("stock_code", "string", False), ("description", "string", True),
         ("quantity", "integer", False), ("invoice_date", "date", False), ("unit_price", "decimal", False), ("customer_id", "string", True),
         ("country", "string", False), ("is_cancellation", "string", False), ("line_amount", "decimal", False), ("source_row", "integer", False)]
    return write_case(cid, f, rows[0], ["One row per invoice line.", "is_cancellation is 'true' when the invoice number starts with C, otherwise 'false'.",
                                        "line_amount is quantity times unit price.", "source_row is the worksheet row of the line."],
                      rows, ["source_row"], "Every invoice line", {"sheet": "Online Retail"}, "Online Retail.xlsx", "sales", "xlsx")


# ---- 8. UK gender pay gap: HR, one row per employer ---------------------------------------------------------------------------
def gpg():
    cid = "uk_gender_pay_gap"
    rows = []
    for r in csv.DictReader(open(CASES / cid / "gender_pay_gap_2023.csv", encoding="utf-8-sig", newline="")):
        rows.append({"source_id": cid, "employer_name": r["EmployerName"], "employer_id": r["EmployerId"], "postcode": r["PostCode"],
                     "company_number": r["CompanyNumber"], "mean_hourly_pay_gap_percent": r["DiffMeanHourlyPercent"],
                     "median_hourly_pay_gap_percent": r["DiffMedianHourlyPercent"], "mean_bonus_gap_percent": r["DiffMeanBonusPercent"],
                     "employer_size": r["EmployerSize"], "submitted_after_deadline": r["SubmittedAfterTheDeadline"].lower(),
                     "due_date": r["DueDate"][:10].replace("/", "-"), "date_submitted": r["DateSubmitted"][:10].replace("/", "-")})
    f = [("source_id", "string", False), ("employer_name", "string", False), ("employer_id", "string", False), ("postcode", "string", True),
         ("company_number", "string", True), ("mean_hourly_pay_gap_percent", "decimal", False), ("median_hourly_pay_gap_percent", "decimal", False),
         ("mean_bonus_gap_percent", "decimal", True), ("employer_size", "string", False), ("submitted_after_deadline", "string", False),
         ("due_date", "date", False), ("date_submitted", "date", False)]
    return write_case(cid, f, rows[0], ["One row per employer.", "submitted_after_deadline is written true or false in lower case.",
                                        "Dates are ISO dates (the time of day is dropped)."],
                      rows, ["employer_id"], "Every employer row", {"file": "gender_pay_gap_2023.csv"}, "gender_pay_gap_2023.csv", "hr", "csv")


def main() -> None:
    cases, scoring = [], {}
    for build in (wb, ons, census_state, census_county, eurostat, cdc, retail, gpg):
        out = build()
        cases.append(out["case"])
        scoring[out["case"]["case_id"]] = out["scoring"]
        print(f"{out['case']['case_id']:24} rows {out['case']['expected_rows']}")
    (ROOT / "benchmark_manifest.json").write_text(json.dumps({"cases": cases, "scoring_cases": scoring}, indent=1), encoding="utf-8")


if __name__ == "__main__":
    main()
