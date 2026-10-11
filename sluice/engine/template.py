"""Step 3b: build the output frame from the template spec and write it into the template workbook."""
from __future__ import annotations

import math
import os
import re
import tempfile
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook

from sluice.engine.context import RunContext
from sluice.engine.diagnosis import EngineError
from sluice.spec.models import ColExpr, TemplateSpec


def _majority(s: pd.Series) -> Any:
    s = s.dropna()
    return s.mode().iloc[0] if not s.empty else pd.NA


def _first(s: pd.Series) -> Any:
    s = s.dropna()
    return s.iloc[0] if not s.empty else pd.NA


def _join(s: pd.Series) -> str:
    return ";".join(sorted({str(v) for v in s.dropna() if str(v) != ""}))


def _sum(s: pd.Series) -> float:
    n = pd.to_numeric(s, errors="coerce")
    present = s.notna() & s.astype(str).str.strip().ne("")
    bad = present & (n.isna() | ~n.map(lambda v: pd.isna(v) or math.isfinite(v)))
    if bad.any():
        raise EngineError("malformed_amount", "Output sum contains non-numeric or non-finite values",
                          stage="transformation", sample=s[bad].head(5).tolist())
    return math.fsum(n.dropna())


AGGS = {
    "first": _first, "majority": _majority, "join": _join, "sum": _sum,
    "last": lambda s: (s.dropna().iloc[-1] if not s.dropna().empty else pd.NA),
    "min": lambda s: pd.to_numeric(s, errors="coerce").min(), "max": lambda s: pd.to_numeric(s, errors="coerce").max(),
    "mean": lambda s: pd.to_numeric(s, errors="coerce").mean(), "count": lambda s: int(s.notna().sum()),
    "nunique": lambda s: int(s.nunique(dropna=True)),
}


def _mask(df: pd.DataFrame, where: dict[str, Any] | None) -> pd.Series:
    mask = pd.Series(True, index=df.index)
    for col, want in (where or {}).items():
        wants = want if isinstance(want, list) else [want]
        mask &= df[col].astype(str).str.strip().str.upper().isin([str(w).upper() for w in wants])
    return mask


def _coerce(s: pd.Series, kind: str | None) -> pd.Series:
    if kind == "date":
        return pd.to_datetime(s, errors="coerce")
    if kind in ("number", "int"):
        n = pd.to_numeric(s, errors="coerce")
        return n.round().astype("Int64") if kind == "int" else n
    if kind == "text":
        return s.map(lambda v: pd.NA if pd.isna(v) else str(v))
    return s


def build_output_frame(df: pd.DataFrame, t: TemplateSpec, ctx: RunContext) -> pd.DataFrame:
    for h, e in t.columns.items():
        for c in [e.col, *(e.where or {})]:
            if c and c not in df.columns:
                raise EngineError("missing_required_column", f"Template column {h!r} needs internal column {c!r}, "
                                  "which no map/shape step produced", stage="transformation",
                                  available=list(df.columns))
    if t.group_by:
        gb = df.groupby(t.group_by, sort=False, dropna=False)
        index = gb.size().index
        out = pd.DataFrame(index=index)
        for h, e in t.columns.items():
            out[h] = _group_col(df, e, t.group_by, index)
        out = out.reset_index(drop=True)
    else:
        out = pd.DataFrame(index=df.index)
        for h, e in t.columns.items():
            if e.col is None:
                out[h] = e.const
            else:
                s = df[e.col].where(_mask(df, e.where))
                out[h] = s if e.default is None else s.fillna(e.default)
        out = out.reset_index(drop=True)
    for h, e in t.columns.items():
        out[h] = _coerce(out[h], e.type)
    if t.sort_by:
        out = out.sort_values(t.sort_by, kind="stable").reset_index(drop=True)
    return out


def _group_col(df: pd.DataFrame, e: ColExpr, by: list[str], index: pd.Index) -> pd.Series:
    if e.col is None:
        return pd.Series([e.const] * len(index), index=index, dtype=object)
    fn = AGGS.get(e.agg)
    if fn is None:
        raise EngineError("processing_failed", f"Unknown aggregation {e.agg!r}; known: {sorted(AGGS)}", stage="transformation")
    sub = df[_mask(df, e.where)]
    res = sub.groupby(by, sort=False, dropna=False)[e.col].agg(fn).reindex(index)
    default = e.default if e.default is not None else (0.0 if e.agg in ("sum", "count", "nunique") else pd.NA)
    return res.astype(object).where(res.notna(), default) if default is not pd.NA else res.astype(object)


def output_name(t: TemplateSpec, out: pd.DataFrame, stem: str, spec_name: str) -> str:
    def repl(m: re.Match[str]) -> str:
        kind, _, arg = m.group(1).partition(":")
        if kind == "stem":
            return stem
        if kind == "name":
            return spec_name
        if kind == "col":
            v = _majority(out[arg]) if arg in out.columns and not out.empty else pd.NA
        elif kind == "yyyymm":
            d = pd.to_datetime(out[arg], errors="coerce").dropna() if arg in out.columns else pd.Series(dtype="datetime64[ns]")
            v = d.dt.strftime("%Y%m").mode().iloc[0] if not d.empty else pd.NA
        else:
            return m.group(0)
        return re.sub(r"[^A-Za-z0-9_.-]+", "_", str(v)).strip("_") if not pd.isna(v) else "UNKNOWN"

    return re.sub(r"\{([^{}]+)\}", repl, t.filename)


def _cell(v: Any) -> Any:
    if v is pd.NA or v is pd.NaT or (isinstance(v, float) and math.isnan(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    return v.item() if hasattr(v, "item") else v


def write_output(out: pd.DataFrame, t: TemplateSpec, template_path: Path | None, dest: Path) -> None:
    """Replace the destination only after the complete file has been written successfully."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=".sluice-", suffix=dest.suffix, dir=dest.parent)
    os.close(fd)
    temp = Path(name)
    try:
        _write_output(out, t, template_path, temp)
        os.replace(temp, dest)
    finally:
        temp.unlink(missing_ok=True)


def _write_output(out: pd.DataFrame, t: TemplateSpec, template_path: Path | None, dest: Path) -> None:
    if dest.suffix.lower() == ".csv":          # flat-file target: the template is just the column list
        dest.parent.mkdir(parents=True, exist_ok=True)
        out.to_csv(dest, index=False, encoding="utf-8")
        return
    if template_path is None:
        wb = Workbook()
        ws = wb.active
        ws.append(list(out.columns))
        header_row = 1
    else:
        wb = load_workbook(template_path)
        ws = wb[t.sheet] if t.sheet else wb.worksheets[0]
        header_row = t.header_row
        headers = [c.value for c in ws[header_row] if c.value is not None]
        if set(headers) != set(out.columns):
            raise EngineError(
                "template_mismatch", "Template headers differ from the spec's template.columns",
                stage="finalization", only_in_template=sorted(set(headers) - set(out.columns)),
                only_in_spec=sorted(set(out.columns) - set(headers)),
            )
        out = out[headers]
        if ws.max_row > header_row:
            ws.delete_rows(header_row + 1, ws.max_row - header_row)
    for r, rec in enumerate(out.itertuples(index=False), start=header_row + 1):
        for c, v in enumerate(rec, start=1):
            ws.cell(row=r, column=c, value=_cell(v))
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    wb.close()
