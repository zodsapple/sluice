"""Step 1: find the right sheet, the header row(s) and the data region."""
from __future__ import annotations

import csv
import io
import re
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

import pandas as pd

from sluice.engine.context import RunContext
from sluice.engine.diagnosis import EngineError
from sluice.spec.models import LocateSpec, MapEntry


def norm(value: Any) -> str:
    """Header normalisation: case/space/punctuation-insensitive."""
    return re.sub(r"[\W_]+", "", str(value).lower())


def expected_headers(entries: dict[str, MapEntry]) -> set[str]:
    out: set[str] = set()
    for src, entry in entries.items():
        out.add(norm(src))
        out.update(norm(a) for a in entry.aliases)
    return out


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NA or str(v).strip() == ""


def _row_score(values: list[Any], expected: set[str], cutoff: float) -> int:
    score = 0
    for v in values:
        if _blank(v):
            continue
        n = norm(v)
        if n in expected or any(SequenceMatcher(None, n, e).ratio() >= cutoff for e in expected):
            score += 1
    return score


def detect_header(raw: pd.DataFrame, expected: set[str], scan: int, cutoff: float) -> tuple[int, int]:
    """Return (row_index, score) of the row that looks most like the expected header."""
    best_idx, best = 0, -1
    for i in range(min(scan, len(raw))):
        s = _row_score(raw.iloc[i].tolist(), expected, cutoff)
        if s > best:
            best_idx, best = i, s
    return best_idx, max(best, 0)


def _read_csv(path: Path, encoding: str | None, delimiter: str | None) -> pd.DataFrame:
    last: Exception | None = None
    for enc in ([encoding] if encoding else ["utf-8-sig", "cp1252", "latin-1"]):
        try:
            text = path.read_text(encoding=enc)
            break
        except UnicodeDecodeError as exc:
            last = exc
    else:
        raise EngineError("input_unreadable", f"Cannot decode {path.name}: {last}", stage="extraction")
    sep = delimiter
    if not sep:
        try:
            sep = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            sep = ","
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=sep))
    width = max((len(r) for r in rows), default=0)
    return pd.DataFrame([r + [None] * (width - len(r)) for r in rows], dtype=object)


def _read_xlsx(path: Path) -> dict[str, pd.DataFrame]:
    """Cell values as published: error cells keep their text (#N/A), blank-looking strings stay strings, "n/a" is not a blank.
    The declared used range is ignored (some exports claim A1)."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: dict[str, pd.DataFrame] = {}
    try:
        for ws in wb.worksheets:
            ws.reset_dimensions()
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            width = max((len(r) for r in rows), default=0)
            out[ws.title] = pd.DataFrame([r + [None] * (width - len(r)) for r in rows], dtype=object)
    finally:
        wb.close()
    return out


def read_workbook(path: Path, locate: LocateSpec) -> dict[str, pd.DataFrame]:
    try:
        if path.suffix.lower() in (".csv", ".tsv", ".tab", ".txt"):
            return {path.stem: _read_csv(path, locate.encoding, locate.delimiter or ("	" if path.suffix.lower() in (".tsv", ".tab") else None))}
        if path.suffix.lower() in (".xlsx", ".xlsm"):
            return _read_xlsx(path)
        engine = "xlrd" if path.suffix.lower() == ".xls" else "openpyxl"
        return pd.read_excel(path, sheet_name=None, header=None, dtype=object, engine=engine,
                             keep_default_na=False, na_values=[""])          # "n/a", "NA", "#N/A" are published tokens, not blanks
    except EngineError:
        raise
    except Exception as exc:
        raise EngineError("input_unreadable", f"Cannot read {path.name}: {exc}", stage="extraction") from exc


def choose_sheet(sheets: dict[str, pd.DataFrame], locate: LocateSpec, entries: dict[str, MapEntry],
                 cutoff: float) -> tuple[str, int]:
    sel = locate.sheet
    names = list(sheets)
    if sel.name:
        hit = next((n for n in names if n.lower() == sel.name.lower()), None)
        if not hit:
            raise EngineError("sheet_not_found", f"Sheet {sel.name!r} not found", stage="extraction",
                              available_sheets=names)
        return hit, 0
    if sel.pattern:
        hit = next((n for n in names if re.search(sel.pattern, n, re.I)), None)
        if not hit:
            raise EngineError("sheet_not_found", f"No sheet matches {sel.pattern!r}", stage="extraction",
                              available_sheets=names)
        return hit, 0
    if sel.strategy == "largest" or len(names) == 1:
        return max(names, key=lambda n: int(sheets[n].notna().any(axis=1).sum())), 0
    exp = expected_headers(entries)
    scored = []
    for n in names:
        _, s = detect_header(sheets[n], exp, locate.max_header_scan, cutoff)
        scored.append((s, int(sheets[n].notna().any(axis=1).sum()), n))
    scored.sort(reverse=True)
    return scored[0][2], scored[0][0]


def flatten_header(raw: pd.DataFrame, rows: list[int]) -> list[str]:
    """Multi-level header -> "top|sub" names. A blank upper cell inherits the group label on its left
    only when a lower-level cell exists (merged group header)."""
    levels = [raw.iloc[r].tolist() for r in rows]
    ncol = len(levels[0])
    filled = [list(lv) for lv in levels]
    for k in range(len(levels) - 1):
        carry = None
        for j in range(ncol):
            v = levels[k][j]
            if not _blank(v):
                carry = v
            elif carry is not None and any(not _blank(levels[m][j]) for m in range(k + 1, len(levels))):
                filled[k][j] = carry
            else:
                carry = None
    names = []
    for j in range(ncol):
        parts = [str(filled[k][j]).strip() for k in range(len(levels)) if not _blank(filled[k][j])]
        names.append("|".join(parts) if parts else f"Unnamed: {j}")
    return names


def build_frame(raw: pd.DataFrame, sheet: str, locate: LocateSpec, entries: dict[str, MapEntry], cutoff: float,
                ctx: RunContext) -> pd.DataFrame:
    """Cut the data region out of a raw sheet and attach lineage (_src_sheet, _src_row)."""
    if locate.synthetic_header is not None:
        if locate.data_start is None:
            raise EngineError("header_not_found", "synthetic_header needs data_start", stage="extraction", sheet=sheet)
        hidx = locate.data_start - 1
        names = [locate.synthetic_header[i] if i < len(locate.synthetic_header) else f"Unnamed: {i}" for i in range(raw.shape[1])]
    elif locate.header_rows:
        hrows = list(locate.header_rows)
        if max(hrows) >= len(raw):
            raise EngineError("header_not_found", f"Header rows {hrows} are beyond the {len(raw)} rows of {sheet!r}",
                              stage="extraction", sheet=sheet)
        hidx = max(hrows)
        names = flatten_header(raw, hrows)
    else:
        if locate.header == "auto":
            hidx, score = detect_header(raw, expected_headers(entries), locate.max_header_scan, cutoff)
            need = min(2, len(entries))
            if score < need:
                preview = [[None if _blank(v) else str(v) for v in raw.iloc[i].tolist()] for i in range(min(5, len(raw)))]
                raise EngineError(
                    "header_not_found",
                    f"No header row in sheet {sheet!r} matches the configured columns "
                    f"(best row had {score} of {len(entries)} matches)",
                    stage="extraction", sheet=sheet, preview_rows=preview,
                )
        else:
            hidx = int(locate.header)
            if hidx >= len(raw):
                raise EngineError("header_not_found", f"Header row {hidx} is beyond the {len(raw)} rows of {sheet!r}",
                                  stage="extraction", sheet=sheet)
        names = []
        for i, v in enumerate(raw.iloc[hidx].tolist()):
            names.append(f"Unnamed: {i}" if _blank(v) else str(v).strip())
    seen: dict[str, int] = {}
    uniq = []
    for base in names:
        seen[base] = seen.get(base, 0) + 1
        uniq.append(base if seen[base] == 1 else f"{base}.{seen[base] - 1}")

    start = locate.data_start if locate.data_start is not None else hidx + 1
    from openpyxl.utils import get_column_letter
    ctx.colrefs[sheet] = {n: get_column_letter(j + 1) for j, n in enumerate(uniq)}
    body = raw.iloc[start:].copy()
    body.columns = uniq
    body.insert(0, "_src_row", body.index + 1)       # 1-based Excel row number
    body.insert(0, "_src_sheet", sheet)
    body[uniq] = body[uniq].apply(lambda s: s.map(lambda v: pd.NA if (v is None or v is pd.NA or (isinstance(v, float) and pd.isna(v)) or v == "") else v))

    blank_rows = body[uniq].isna().all(axis=1)
    if locate.stop_at_blank_row and blank_rows.any():
        body = body.iloc[: blank_rows.to_numpy().argmax()]
    else:
        body = body[~blank_rows]
    keep = [c for c in uniq if not (c.startswith("Unnamed: ") and body[c].isna().all())]
    body = body[["_src_sheet", "_src_row", *keep]]

    for pat in locate.skip_rows_matching:
        rx = re.compile(pat, re.I)

        def first_cell(r: pd.Series) -> str:
            for v in r:
                if not pd.isna(v):
                    return str(v)
            return ""

        label = body[keep].apply(first_cell, axis=1)
        hit = label.map(lambda t: bool(rx.search(t)))
        for _, row in body[hit].iterrows():
            ctx.issue("skipped_total_row", f"Row matches skip pattern {pat!r}", severity="info", action="skipped",
                      stage="extraction", sheet=sheet, row=row["_src_row"], raw=first_cell(row[keep]))
        body = body[~hit]
    return body.reset_index(drop=True)


def load_source(path: Path, spec: Any, ctx: RunContext) -> pd.DataFrame:
    """Primary source of a pipeline spec (kept for the single-source call path)."""
    sheets = read_workbook(path, spec.locate)
    ctx.workbook_sheets = sheets
    sheet, _ = choose_sheet(sheets, spec.locate, spec.map, spec.map_cutoff)
    ctx.sheet_name = sheet
    return build_frame(sheets[sheet], sheet, spec.locate, spec.map, spec.map_cutoff, ctx)
